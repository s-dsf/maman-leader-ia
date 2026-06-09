import os
import hashlib
from pathlib import Path
import chromadb
from chromadb.utils import embedding_functions

DOCS_DIR = Path("documents")
DOCS_DIR.mkdir(exist_ok=True)
CHROMA_DIR = Path("chroma_db")

client_chroma = chromadb.PersistentClient(path=str(CHROMA_DIR))
ef = embedding_functions.DefaultEmbeddingFunction()
collection = client_chroma.get_or_create_collection(
    name="maison_edition",
    embedding_function=ef
)

def extraire_texte(fichier: Path) -> str:
    ext = fichier.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(fichier))
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(str(fichier))
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext in [".xlsx", ".xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(str(fichier))
            texte = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    ligne = " | ".join(str(c) for c in row if c)
                    if ligne:
                        texte.append(ligne)
            return "\n".join(texte)
        elif ext == ".txt":
            return fichier.read_text(encoding="utf-8", errors="ignore")
        else:
            return ""
    except Exception as e:
        print(f"Erreur lecture {fichier.name}: {e}")
        return ""

def decouper_texte(texte: str, taille=800, chevauchement=100) -> list:
    mots = texte.split()
    chunks = []
    i = 0
    while i < len(mots):
        chunk = " ".join(mots[i:i+taille])
        if chunk.strip():
            chunks.append(chunk)
        i += taille - chevauchement
    return chunks

def indexer_documents():
    fichiers = [f for f in DOCS_DIR.glob("**/*")
                if f.is_file() and f.suffix.lower()
                in [".pdf", ".docx", ".txt", ".xlsx", ".xls"]]
    if not fichiers:
        return 0, []
    indexed = []
    for fichier in fichiers:
        texte = extraire_texte(fichier)
        if not texte.strip():
            continue
        chunks = decouper_texte(texte)
        ids, docs, metas = [], [], []
        for i, chunk in enumerate(chunks):
            chunk_id = hashlib.md5(
                f"{fichier.name}_{i}_{chunk[:50]}".encode()
            ).hexdigest()
            ids.append(chunk_id)
            docs.append(chunk)
            metas.append({"source": fichier.name, "chunk": i})
        if ids:
            collection.upsert(documents=docs, ids=ids, metadatas=metas)
            indexed.append(fichier.name)
    return len(indexed), indexed

def rechercher(query: str, n=4) -> str:
    try:
        total = collection.count()
        if total == 0:
            return ""
        n = min(n, total)
        results = collection.query(query_texts=[query], n_results=n)
        if not results["documents"] or not results["documents"][0]:
            return ""
        sources = []
        for doc, meta in zip(
            results["documents"][0], results["metadatas"][0]
        ):
            sources.append(f"[Source: {meta['source']}]\n{doc}")
        return "\n\n---\n\n".join(sources)
    except Exception:
        return ""

def stats_base():
    try:
        return collection.count()
    except:
        return 0